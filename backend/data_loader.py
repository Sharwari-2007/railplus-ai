import os
import math
import openpyxl

def haversine_distance(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Calculates Haversine distance in km between two lat/lon coordinates"""
    R = 6371.0  # Earth radius in kilometers
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = math.sin(dlat / 2)**2 + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(dlon / 2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c

class WeatherDataLoader:
    def __init__(self, excel_path: str = None):
        if excel_path is None:
            base_dir = os.path.dirname(__file__)
            excel_path = os.path.join(base_dir, "data", "india_weather_rainfall_data.xlsx")
        
        self.excel_path = excel_path
        self.records = []
        self.station_coords = {}
        self.unique_stations = []
        self.load_data()

    def load_data(self):
        if not os.path.exists(self.excel_path):
            from generate_dataset import generate_weather_dataset
            generate_weather_dataset()

        print(f"Loading weather dataset from {self.excel_path} via openpyxl...")
        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        ws = wb.active

        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        
        records = []
        unique_stations_set = set()
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            records.append(row_dict)
            st_name = row_dict["station_name"]
            
            if st_name not in unique_stations_set:
                unique_stations_set.add(st_name)
                self.station_coords[st_name] = {
                    "lat": float(row_dict["latitude"]),
                    "lon": float(row_dict["longitude"]),
                    "state": row_dict["state"],
                    "district": row_dict["district"],
                    "elevation": float(row_dict["elevation"])
                }

        self.records = records
        self.unique_stations = list(unique_stations_set)
        wb.close()
        print(f"Weather dataset loaded: {len(self.records)} records across {len(self.unique_stations)} stations.")

    def get_nearest_weather_station(self, lat: float, lon: float) -> str:
        """Weather Proximity Index: Maps any track/station node to nearest met station using Haversine lookup"""
        best_station = self.unique_stations[0]
        min_dist = float("inf")

        for st_name, coords in self.station_coords.items():
            dist = haversine_distance(lat, lon, coords["lat"], coords["lon"])
            if dist < min_dist:
                min_dist = dist
                best_station = st_name

        return best_station

    def get_weather_for_node(self, lat: float, lon: float, date_str: str = "2025-07-15") -> dict:
        nearest_st = self.get_nearest_weather_station(lat, lon)
        
        # Search records for station & date
        for r in self.records:
            if r["station_name"] == nearest_st and r["date_of_record"] == date_str:
                res = dict(r)
                res["nearest_met_station"] = nearest_st
                return res

        # Fallback to first station record if date missing
        for r in self.records:
            if r["station_name"] == nearest_st:
                res = dict(r)
                res["nearest_met_station"] = nearest_st
                return res

        # Default fallback
        return {
            "station_name": nearest_st,
            "nearest_met_station": nearest_st,
            "rainfall": 0.0,
            "min_temp": 20.0,
            "max_temp": 32.0,
            "avg_temp": 26.0,
            "wind_speed": 15.0,
            "air_pressure": 1012.0,
            "season": "Monsoon"
        }

    def get_historic_days(self) -> list:
        return [
            {"id": "monsoon_peak", "label": "Severe Monsoon Hazard (July 15, 2025)", "date": "2025-07-15", "type": "Monsoon"},
            {"id": "winter_fog", "label": "Winter Dense Fog Wave (January 10, 2025)", "date": "2025-01-10", "type": "Fog"},
            {"id": "summer_storm", "label": "Summer High-Wind Storm (May 20, 2025)", "date": "2025-05-20", "type": "Wind"},
            {"id": "clear_sky", "label": "Optimal Operational Day (October 12, 2025)", "date": "2025-10-12", "type": "Normal"}
        ]

# Global singleton instance
weather_loader = WeatherDataLoader()
